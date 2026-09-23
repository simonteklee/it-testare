"""Textextrahering ur många filtyper + chunkning + webbsidehämtning."""
from __future__ import annotations

import re
import subprocess
import tempfile
from pathlib import Path

import httpx

HEADERS = {"User-Agent": "TestARN/0.1 (privat studieverktyg)"}

OFFICE_EXTS = {".docx", ".doc", ".odt", ".rtf", ".pptx", ".ppt", ".odp",
               ".xlsx", ".xls", ".ods", ".csv", ".txt"}
HTML_EXTS = {".html", ".htm", ".xhtml"}


def html_to_text(html: str) -> str:
    html = re.sub(r"(?is)<(script|style|nav|footer|header|aside)[^>]*>.*?</\1>", " ", html)
    html = re.sub(r"(?is)<br\s*/?>", "\n", html)
    html = re.sub(r"(?is)</(p|div|li|h[1-6]|tr)>", "\n", html)
    text = re.sub(r"(?s)<[^>]+>", " ", html)
    import html as _h
    text = _h.unescape(text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


def _soffice_to_text(path: Path) -> str:
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            ["soffice", "--headless", "-env:UserInstallation=file:///tmp/loprofile-simi",
             "--convert-to", "txt:Text (encoded):UTF8", str(path), "--outdir", tmp],
            check=False, capture_output=True, timeout=120,
        )
        cands = list(Path(tmp).glob("*.txt"))
        if not cands:
            raise RuntimeError("kunde inte konvertera filen (soffice)")
        return cands[0].read_text(encoding="utf-8", errors="ignore")


def _pdf_text(path: Path) -> str:
    try:
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as t:
            out = t.name
        subprocess.run(["pdftotext", "-layout", str(path), out], check=True,
                       capture_output=True, timeout=120)
        return Path(out).read_text(encoding="utf-8", errors="ignore")
    except Exception:
        pass
    try:  # portabelt (Windows/macOS utan pdftotext)
        from pypdf import PdfReader
        pages = PdfReader(str(path)).pages
        return "\n".join((p.extract_text() or "") for p in pages)
    except Exception as e:  # noqa: BLE001
        raise RuntimeError(f"kunde inte läsa PDF: {e}")


def _office_via_libs(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".docx":
        import docx
        return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)
    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(str(path))
        out = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    out.append(shape.text_frame.text)
        return "\n".join(out)
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                out.append("\t".join("" if c is None else str(c) for c in row))
        return "\n".join(out)
    raise RuntimeError("formatet kräver LibreOffice")


def _office_text(path: Path) -> str:
    try:
        return _soffice_to_text(path)
    except Exception:
        return _office_via_libs(path)


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _pdf_text(path)
    if ext in HTML_EXTS:
        return html_to_text(path.read_text(encoding="utf-8", errors="ignore"))
    if ext in OFFICE_EXTS and ext not in {".txt", ".csv"}:
        return _office_text(path)
    # txt, csv, md, json, kod, m.m. → läs som text
    try:
        return path.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return _office_text(path)


def fetch_url(url: str) -> str:
    with httpx.Client(timeout=60.0, follow_redirects=True, headers=HEADERS) as c:
        r = c.get(url)
        r.raise_for_status()
        ctype = r.headers.get("content-type", "")
        if "html" in ctype or url.lower().endswith((".html", ".htm")):
            return html_to_text(r.text)
        return r.text


def chunk(text: str, size: int = 900) -> list[str]:
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks, cur = [], ""
    for p in paras:
        if len(cur) + len(p) + 1 <= size:
            cur = (cur + "\n" + p).strip()
        else:
            if cur:
                chunks.append(cur)
            cur = p if len(p) <= size else p[:size]
    if cur:
        chunks.append(cur)
    return [c for c in chunks if len(c) > 60]
